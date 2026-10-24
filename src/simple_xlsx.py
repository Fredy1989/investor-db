"""Jednoduché Excely pro člověka: jeden řádek = jeden investor, jen to, co chce zadání.

Zadání: u každého investora sektor, do čeho investuje a v jaké výši, celkový kapitál a odkud informace jsou;
u každého záznamu zdroj a datum, které dokládají, že subjekt investuje.

  python src/simple_xlsx.py databaze     → plocha / „Databáze investorů.xlsx"  (zařazení + vyřazení)
  python src/simple_xlsx.py kontrola     → plocha / „Ruční kontrola.xlsx"      (vylosovaný vzorek, políčko u každého údaje)
  python src/simple_xlsx.py import       ← „Ruční kontrola.xlsx" → data/manual_check.csv (pro src/measure.py)

Každá hodnota je odkaz na svůj zdroj (klik otevře stránku). Plné záznamy se citacemi zůstávají v data/records/.
"""
import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from make_manual_sheet import PINNED_SAMPLE
from records import ROOT, load_records

DESKTOP = Path.home() / "Desktop"
DB_XLSX = DESKTOP / "Databáze investorů.xlsx"
CHECK_XLSX = DESKTOP / "Ruční kontrola.xlsx"
CHECK2_XLSX = DESKTOP / "Ruční kontrola 2.xlsx"
# Vzorek 2: vylosováno 2026-10-24 (random.Random(20261024).sample(21 zařazených z kol 2–3, 6)) – záznamy, kterých se
# netýkalo žádné review navíc (vzorek 1 náhodou prošel i review Mac Clauda, viz reviews/mac-claude-2.md).
SAMPLE2 = ["arx", "bhm", "flyingfounders", "hlavenka", "jic", "soulmates"]
MANUAL_CSV = ROOT / "data" / "manual_check.csv"

TYPE_CZ = {"vc": "VC fond", "pe": "private equity", "cvc": "korporátní VC", "public_vc": "státní VC",
           "family_office": "family office", "angel": "anděl", "angel_syndicate": "andělský syndikát",
           "private_investor": "soukromý investor"}
SECTOR_CZ = {"fintech": "fintech", "saas_b2b": "B2B software", "deeptech": "deep tech", "ai": "AI",
             "health": "zdraví", "energy_climate": "energetika a klima", "consumer": "spotřebitelé",
             "ecommerce": "e-commerce", "defence": "obrana", "mobility": "mobilita", "industry": "průmysl",
             "generalist": "bez omezení oboru"}
STAGE_CZ = {"pre_seed": "pre-seed", "seed": "seed", "series_a": "série A", "series_b_plus": "série B+",
            "growth": "růstová fáze", "buyout": "buyout"}
BASIS_CZ = {"latest_fund_size": "poslední fond", "target_fund_size": "cíl fondu", "aum": "spravovaná aktiva",
            "total_across_funds": "součet fondů", "own_balance_sheet": "vlastní kapitál",
            "portfolio_value": "hodnota portfolia", "previous_fund_size": "předchozí fond",
            "total_assets": "aktiva celé skupiny"}
RULE_CZ = {"E1": "crowdfunding / platforma, sám neinvestuje", "E2": "akcelerátor bez investičního vozidla",
           "E3": "jen dluhové financování", "E4": "síť / asociace / poradce, sám neinvestuje",
           "E5": "investuje jen do burzovních akcií nebo realit", "E6": "žádná doložená investice za 3 roky",
           "E7": "strategický kupec, ne investor", "E8": "duplicita",
           "I1": "nepodařilo se doložit, že subjekt existuje", "I2": "žádná doložená investice",
           "I4": "neinvestuje přímo do firem", "out_of_scope": "mimo rozsah vzorku (bez českého týmu / PE)"}
NONE = "neuvedeno"
CHECK = ["sedí", "nesedí", "nejde ověřit"]
HEAD = PatternFill("solid", fgColor="1F3864")
ALT = PatternFill("solid", fgColor="EEF3FA")
YOU = PatternFill("solid", fgColor="FFF4CC")


def eur(x: float | None) -> str:
    if x is None:
        return ""
    def num(v: float) -> str:
        return f"{v:.1f}".replace(".", ",").removesuffix(",0")
    if x >= 1e9:
        return f"{num(x / 1e9)} mld. EUR"
    if x >= 1e6:
        return f"{num(x / 1e6)} mil. EUR"
    return f"{x / 1e3:.0f} tis. EUR"


def fact(rec: dict, key: str) -> tuple[str, str | None]:
    """(text pro člověka, URL zdroje) pro sektor / fáze / tiket / kapitál."""
    f = rec.get(key) or {}
    v, url = f.get("value"), f.get("source_url")
    if key == "sectors":
        specific = [s for s in (v or []) if s != "generalist"]
        text = ", ".join(SECTOR_CZ.get(s, s) for s in (specific or v or []))
    elif key == "stages":
        text = ", ".join(STAGE_CZ.get(s, s) for s in (v or []))
    elif key == "ticket":
        lo, hi = (v or {}).get("min_eur"), (v or {}).get("max_eur")
        text = f"{eur(lo)} – {eur(hi)}" if lo and hi else f"až {eur(hi)}" if hi else f"od {eur(lo)}" if lo else ""
    else:
        amount, basis = (v or {}).get("eur"), (v or {}).get("basis")
        text = f"{eur(amount)} ({BASIS_CZ.get(basis, basis)})" if amount else ""
    return (text or NONE), (url if text else None)


def deal(rec: dict) -> tuple[str, str | None]:
    deals = [d for d in rec.get("investment_evidence") or [] if d.get("date")]
    if not deals:
        return NONE, None
    d = max(deals, key=lambda d: d["date"])
    y, m, dd = (d["date"] + "--").split("-")[:3]
    date = f"{int(dd)}. {int(m)}. {y}" if dd else f"{int(m)}/{y}" if m else y
    return f"{d['company']} ({date})", d.get("source_url")


def put(ws, row: int, col: int, text: str, url: str | None = None, fill=None) -> None:
    cell = ws.cell(row, col, text)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    if url:
        cell.hyperlink = url
        cell.font = Font(color="0563C1", underline="single")
    elif text == NONE:
        cell.font = Font(color="999999", italic=True)
    if fill:
        cell.fill = fill


COLUMN_HELP = {  # záhlaví → (co znamená, příklad); použito v listu Vysvětlivky i jako komentář u záhlaví
    "Investor": ("Název investora. Klik otevře jeho web.", "Credo Ventures"),
    "Typ": ("Druh investora, tabulka typů je v listu Vysvětlivky.", "VC fond, private equity, anděl"),
    "Sektor": ("V jakých oborech investor investuje (podle toho, co sám deklaruje).",
               "AI, fintech, zdraví; „bez omezení oboru“ = investuje do čehokoli"),
    "Do čeho investuje (fáze)": ("V jakém stadiu vývoje jsou firmy, do kterých investuje — tabulka fází je v listu Vysvětlivky.",
                                 "pre-seed, seed = velmi mladé firmy"),
    "Fáze": ("V jakém stadiu vývoje jsou firmy, do kterých investuje — tabulka fází je v listu Vysvětlivky.",
             "pre-seed, seed = velmi mladé firmy"),
    "Typická výše investice": ("Kolik obvykle vloží do jedné firmy (tzv. tiket).", "250 tis. – 1 mil. EUR"),
    "Výše investice": ("Kolik obvykle vloží do jedné firmy (tzv. tiket).", "250 tis. – 1 mil. EUR"),
    "Celkový kapitál": ("Kolik peněz má investor celkem; v závorce co číslo znamená (tabulka v listu Vysvětlivky).",
                        "75,9 mil. EUR (poslední fond)"),
    "Doklad, že investuje (firma, datum)": ("Poslední veřejně doložená investice za poslední 3 roky. "
                                            "Klik otevře článek nebo tiskovou zprávu.", "Modeinspect (15. 1. 2026)"),
    "Doklad, že investuje": ("Poslední veřejně doložená investice za poslední 3 roky. "
                             "Klik otevře článek nebo tiskovou zprávu.", "Modeinspect (15. 1. 2026)"),
    "✔": ("Tvoje odpověď: sedí / nesedí / nejde ověřit.", "sedí"),
}
STAGE_HELP = [
    ("pre-seed", "jen nápad nebo první prototyp, často ještě bez firmy", "desítky až stovky tisíc EUR"),
    ("seed", "první produkt, první zákazníci", "stovky tisíc až ~2 mil. EUR"),
    ("série A", "produkt funguje, firma roste a potřebuje peníze na škálování", "jednotky milionů EUR"),
    ("série B+", "další růstová kola, expanze do zahraničí", "desítky milionů EUR"),
    ("růstová fáze", "zavedená firma, investor kupuje menšinový podíl", "desítky milionů EUR"),
    ("buyout", "investor kupuje celou firmu nebo většinu (typické pro private equity)", "desítky až stovky mil. EUR"),
]
BASIS_HELP = [
    ("poslední fond", "velikost posledního uzavřeného fondu"),
    ("cíl fondu", "kolik fond chce vybrat — uzavření zatím není doložené"),
    ("předchozí fond", "velikost staršího fondu (aktuální není veřejná)"),
    ("součet fondů", "všechny fondy investora dohromady"),
    ("spravovaná aktiva", "vše, co investor spravuje (AUM), i mimo startupy"),
    ("hodnota portfolia", "současná hodnota firem, ve kterých má podíl — ne volné peníze"),
    ("vlastní kapitál", "investor investuje vlastní peníze (bez cizích investorů)"),
    ("aktiva celé skupiny", "majetek celé skupiny vč. nemovitostí — ne jen peníze na investice"),
]
TYPE_HELP = [
    ("VC fond", "fond rizikového kapitálu, investuje do startupů"),
    ("private equity", "kupuje podíly nebo celé zralé firmy"),
    ("korporátní VC", "investiční rameno firmy nebo banky"),
    ("family office", "spravuje a investuje majetek jedné rodiny (a tak se sám označuje)"),
    ("soukromý investor", "investiční skupina podnikatele, investuje vlastní peníze"),
    ("anděl", "fyzická osoba investující vlastní peníze do startupů"),
    ("andělský syndikát", "skupina andělů investující přes společnou firmu"),
]


def header(ws, titles: list[tuple[str, int]]) -> None:
    for i, (title, width) in enumerate(titles, 1):
        c = ws.cell(1, i, title)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = width
        if title in COLUMN_HELP:  # vysvětlení po najetí myší na záhlaví
            meaning, example = COLUMN_HELP[title]
            c.comment = Comment(f"{meaning}\nPříklad: {example}", "investor-db", width=320, height=110)
    ws.freeze_panes = "B2"


def glossary(wb, columns: list[str]) -> None:
    """List „Vysvětlivky“: co znamená každý sloupec + tabulky fází, typů a druhů kapitálu."""
    ws = wb.create_sheet("Vysvětlivky", 1)

    def section(title: str, head: tuple, rows: list) -> None:
        ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=13)
        ws.append(list(head))
        for c in ws[ws.max_row]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
        for r in rows:
            ws.append(list(r))
            for c in ws[ws.max_row]:
                c.alignment = Alignment(wrap_text=True, vertical="top")

    ws.append(["Vysvětlivky"])
    ws.cell(1, 1).font = Font(bold=True, size=15)
    section("Co znamená který sloupec", ("Sloupec", "Co znamená", "Příklad"),
            [(c, *COLUMN_HELP[c]) for c in dict.fromkeys(columns) if c in COLUMN_HELP])
    section("Fáze — v jakém stadiu jsou firmy, do kterých investor investuje",
            ("Fáze", "Co to znamená", "Typická investice"), STAGE_HELP)
    section("Celkový kapitál — co znamená údaj v závorce", ("V závorce", "Co číslo znamená", ""), BASIS_HELP)
    section("Typy investorů", ("Typ", "Co to je", ""), TYPE_HELP)
    ws.append([])
    ws.append(["„neuvedeno“ = investor údaj veřejně neuvádí; neodhadujeme ho. Částky v jiných měnách jsou přepočtené "
               "na EUR kurzem ECB k datu zdroje."])
    for col, width in zip("ABC", (26, 70, 40)):
        ws.column_dimensions[col].width = width


DB_COLUMNS = [("Investor", 28), ("Typ", 16), ("Sektor", 28), ("Do čeho investuje (fáze)", 24),
              ("Typická výše investice", 20), ("Celkový kapitál", 24), ("Doklad, že investuje (firma, datum)", 32)]


def databaze() -> None:
    recs = sorted(load_records(), key=lambda r: r["name"].lower())
    wb = Workbook()
    ws = wb.active
    ws.title = "Investoři"
    header(ws, DB_COLUMNS)
    row = 1
    for r in (r for r in recs if r["decision"] == "include"):
        row += 1
        fill = ALT if row % 2 else None
        put(ws, row, 1, r["name"], r.get("website"), fill)
        put(ws, row, 2, TYPE_CZ.get(r.get("investor_type"), r.get("investor_type") or ""), None, fill)
        for col, key in enumerate(["sectors", "stages", "ticket", "capital"], 3):
            put(ws, row, col, *fact(r, key), fill)
        put(ws, row, 7, *deal(r), fill)

    ex = wb.create_sheet("Vyřazení")
    header(ex, [("Subjekt", 30), ("Proč není v databázi", 50), ("Podrobnosti", 90)])
    for i, r in enumerate((r for r in recs if r["decision"] != "include"), 2):
        rule = r.get("exclusion_rule") or ""
        put(ex, i, 1, r["name"])
        put(ex, i, 2, RULE_CZ.get(rule.split()[0] if rule else "", rule))
        put(ex, i, 3, r.get("decision_reason") or "")

    glossary(wb, [title for title, _ in DB_COLUMNS])
    info = wb.create_sheet("Jak číst")
    for line in [
        "Jeden řádek = jeden investor. Klikni na hodnotu a otevře se zdroj, ze kterého pochází.",
        "Co znamená který sloupec, fáze (pre-seed, seed…) a údaje u kapitálu: list „Vysvětlivky“ "
        "(nebo najeď myší na záhlaví sloupce).",
        "„neuvedeno“ = investor údaj veřejně neuvádí. Neodhadujeme ho.",
        "Částky v jiných měnách přepočteny na EUR kurzem ECB k datu zdroje.",
        "Celkový kapitál: v závorce, co číslo znamená (poslední fond / cíl fondu / spravovaná aktiva…).",
        "Doklad: poslední veřejně doložená investice za poslední 3 roky; klik otevře článek nebo tiskovou zprávu.",
        "Pravidla zařazení: PLAN.md. Plné záznamy s citacemi: data/records/.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 110
    wb.save(DB_XLSX)
    print(f"uloženo: {DB_XLSX}")


CHECK_FIELDS = [("Doklad, že investuje", None), ("Sektor", "sectors"), ("Fáze", "stages"),
                ("Výše investice", "ticket"), ("Celkový kapitál", "capital")]


def kontrola(sample: list[str] = PINNED_SAMPLE, path: Path = CHECK_XLSX, disputed_sheet: bool = True) -> None:
    recs = {r["id"]: r for r in load_records()}
    wb = Workbook()
    ws = wb.active
    ws.title = "Kontrola"
    cols = [("Investor", 24)]
    for title, _ in CHECK_FIELDS:
        cols += [(title, 26), ("✔", 11)]
    header(ws, cols + [("Poznámka", 30)])
    dv = DataValidation(type="list", formula1='"' + ",".join(CHECK) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    done = saved_answers()
    for row, rid in enumerate(sample, 2):
        r = recs[rid]
        put(ws, row, 1, r["name"])
        for i, (_, key) in enumerate(CHECK_FIELDS):
            put(ws, row, 2 + 2 * i, *(deal(r) if key is None else fact(r, key)))
            put(ws, row, 3 + 2 * i, done.get((rid, key or "decision"), ("", ""))[0], None, YOU)
            dv.add(ws.cell(row, 3 + 2 * i))
        put(ws, row, 12, next((c for (i_, _), (_, c) in done.items() if i_ == rid and c), ""), None, YOU)
        ws.row_dimensions[row].height = 48

    sp = wb.create_sheet("Sporné") if disputed_sheet else None
    disputed = []
    if sp:
        header(sp, [("Subjekt", 26), ("Naše rozhodnutí", 34), ("Proč", 70), ("Souhlasíš?", 13), ("Poznámka", 30)])
        sdv = DataValidation(type="list", formula1='"' + ",".join(CHECK) + '"', allow_blank=True)
        sp.add_data_validation(sdv)
        with open(MANUAL_CSV, encoding="utf-8-sig", newline="") as f:
            disputed = list(dict.fromkeys(r["id"] for r in csv.DictReader(f) if r["group"] == "2_disputed"))
    for row, rid in enumerate(disputed, 2):
        r = recs[rid]
        rule = r.get("exclusion_rule") or ""
        verdict = ("zařazen jako " + TYPE_CZ.get(r.get("investor_type"), "?")) if r["decision"] == "include" \
            else RULE_CZ.get(rule.split()[0] if rule else "", rule)
        put(sp, row, 1, r["name"])
        put(sp, row, 2, verdict, deal(r)[1])
        put(sp, row, 3, r.get("decision_reason") or "")
        answer, note = done.get((rid, "decision"), ("", ""))
        put(sp, row, 4, answer, None, YOU)
        sdv.add(sp.cell(row, 4))
        put(sp, row, 5, note, None, YOU)

    glossary(wb, [title for title, _ in cols])
    info = wb.create_sheet("Návod")
    for line in [
        "List „Kontrola“: 6 náhodně vylosovaných investorů, u každého 5 údajů, které chce zadání.",
        "Co znamená který sloupec (např. Fáze = v jakém stadiu jsou firmy, do kterých investor investuje): "
        "list „Vysvětlivky“, nebo najeď myší na záhlaví sloupce.",
        "Klikni na hodnotu (otevře se zdroj) a do žlutého políčka ✔ vedle ní vyber:",
        "   sedí = zdroj údaj potvrzuje (u „neuvedeno“: ani ty ho za 1–2 minuty nenajdeš)",
        "   nesedí = zdroj říká něco jiného (u „neuvedeno“: údaj veřejně existuje) → do Poznámky napiš co",
        "   nejde ověřit = stránka nejde otevřít nebo tam údaj není",
        "Částky: stačí, když sedí částka a měna ve zdroji; přepočet na EUR dělá kód.",
        *(["List „Sporné“: hraniční rozhodnutí — souhlasíš s tím, jak jsme je zařadili?"] if disputed_sheet else []),
        "Hotovo → ulož a řekni Claudovi.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 110
    wb.save(path)
    print(f"uloženo: {path}")


def add_sample2_rows() -> None:
    """Přidá do data/manual_check.csv řádky vzorku 2 (skupina 4_random_sample_2), pokud tam ještě nejsou."""
    with open(MANUAL_CSV, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    if not any(r["group"] == "4_random_sample_2" for r in rows):
        recs = {r["id"]: r for r in load_records()}
        rows += sample2_rows(recs, list(rows[0].keys()))
    # verdikty slepého AI ověřovatele (fáze 2–3) pro měření shody AI × člověk; opravená pole se ze shody vynechají
    with open(ROOT / "data" / "corrections.json", encoding="utf-8") as f:
        corrected = {(c["id"], c["path"].split(".")[0]) for c in json.load(f)}
    for r in rows:
        if r["group"] != "4_random_sample_2":
            continue
        vpath = ROOT / "data" / "verification_v2" / f"{r['id']}.csv"
        with open(vpath, encoding="utf-8-sig", newline="") as f:
            blind = {v["claim"].strip(): v["verdict"].strip().lower() for v in csv.DictReader(f)}
        r["agent_b_verdict"] = blind.get(r["field"], "")
        r["corrected_after_b"] = "yes" if (r["id"], r["field"]) in corrected else ""
    with open(MANUAL_CSV, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def sample2_rows(recs: dict, columns: list[str]) -> list[dict]:
    from make_manual_sheet import fact_row
    rows = []
    for rid in SAMPLE2:
        for field in ["decision", "sectors", "stages", "ticket", "capital", "investment_evidence"]:
            fr = fact_row(recs[rid], field)
            rows.append({**{k: "" for k in columns}, "group": "4_random_sample_2", "id": rid,
                         "name": recs[rid]["name"], "field": field, "value": fr["value"],
                         "source_url": fr["source_url"], "quote": fr["quote"], "rationale": fr.get("rationale", "")})
    return rows


ANSWER_OF = {"ok": "sedí", "null_ok": "sedí", "wrong": "nesedí", "null_missed": "nesedí",
             "unverifiable": "nejde ověřit"}


def saved_answers() -> dict:
    """Už zadané odpovědi z data/manual_check.csv – při novém vytvoření souboru se nesmí ztratit."""
    out = {}
    with open(MANUAL_CSV, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            if r.get("human_verdict"):
                out[(r["id"], r["field"])] = (ANSWER_OF.get(r["human_verdict"], ""), r.get("human_comment", ""))
    return out


def to_verdict(answer: str, empty: bool) -> str:
    a = (answer or "").strip().lower()
    if not a:
        return ""
    if a == "nejde ověřit":
        return "unverifiable"
    if a == "sedí":
        return "null_ok" if empty else "ok"
    if a == "nesedí":
        return "null_missed" if empty else "wrong"
    raise ValueError(f"neznámá odpověď: {answer}")


def import_(path: Path = CHECK_XLSX) -> None:
    wb = load_workbook(path, data_only=True)
    recs = {r["id"]: r for r in load_records()}
    by_name = {r["name"]: rid for rid, r in recs.items()}
    answers = {}  # (id, field) -> (verdict, comment)
    for row in wb["Kontrola"].iter_rows(min_row=2, values_only=True):
        rid = by_name.get(row[0])
        if not rid:
            continue
        for i, (_, key) in enumerate(CHECK_FIELDS):
            value, answer = row[1 + 2 * i], row[2 + 2 * i]
            v = to_verdict(answer, value == NONE)
            for field in (["decision", "investment_evidence"] if key is None else [key]):
                answers[(rid, field)] = (v, row[11] or "")
    for row in (wb["Sporné"].iter_rows(min_row=2, values_only=True) if "Sporné" in wb.sheetnames else []):
        rid = by_name.get(row[0])
        if rid:
            v = to_verdict(row[3], False)
            for field in ("decision", "investor_type"):
                answers[(rid, field)] = (v, row[4] or "")
    with open(MANUAL_CSV, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    n = 0
    for r in rows:
        if (r["id"], r["field"]) in answers:
            r["human_verdict"], r["human_comment"] = answers[(r["id"], r["field"])]
            n += bool(r["human_verdict"])
    with open(MANUAL_CSV, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    print(f"načteno {n} verdiktů do {MANUAL_CSV}")


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    {"databaze": databaze, "kontrola": kontrola, "import": import_,
     "kontrola2": lambda: (add_sample2_rows(), kontrola(SAMPLE2, CHECK2_XLSX, disputed_sheet=False)),
     "import2": lambda: import_(CHECK2_XLSX)}.get(cmd, lambda: sys.exit(__doc__))()
