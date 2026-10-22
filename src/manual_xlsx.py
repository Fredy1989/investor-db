"""Převod listu pro ruční kontrolu mezi CSV (zdroj pravdy v repu) a čitelným Excelem pro člověka.

  python src/manual_xlsx.py export [cesta.xlsx]   data/manual_check.csv → Excel (výchozí: plocha)
  python src/manual_xlsx.py import [cesta.xlsx]   vyplněné verdikty z Excelu → data/manual_check.csv

Excel: údaj hned vedle jména, investoři barevně odlišení, verdikt z rozbalovací nabídky, klikací odkazy.
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from records import ROOT

CSV_PATH = ROOT / "data" / "manual_check.csv"
DEFAULT_XLSX = Path.home() / "Desktop" / "investor-db - rucni kontrola.xlsx"
COLUMNS = [  # (sloupec v CSV, záhlaví v Excelu, šířka)
    ("id", "id", 12), ("name", "investor", 26), ("field", "údaj", 18), ("value", "hodnota v databázi", 45),
    ("source_url", "zdroj (klikni)", 40), ("quote", "citace ze zdroje (najdi Ctrl+F)", 50),
    ("human_verdict", "TVŮJ VERDIKT", 16), ("human_comment", "tvůj komentář", 35),
    ("rationale", "zdůvodnění AI agenta (není ze zdroje)", 40), ("group", "skupina", 18),
    ("agent_b_verdict", "verdikt AI (nečíst předem)", 14), ("corrected_after_b", "opraveno po AI", 10),
]
VERDICTS = ["ok", "wrong", "unverifiable", "null_ok", "null_missed"]
BANDS = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="EEF3FA")]
YOUR_FILL = PatternFill("solid", fgColor="FFF4CC")


def export(xlsx: Path) -> None:
    with open(CSV_PATH, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    wb = Workbook()
    ws = wb.active
    ws.title = "Kontrola"
    ws.append([h for _, h, _ in COLUMNS])
    for i, (_, _, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
        ws.cell(1, i).font = Font(bold=True)
    band, last_id = 0, None
    for r in rows:
        if r["id"] != last_id:
            band, last_id = 1 - band, r["id"]
        ws.append([r.get(col, "") for col, _, _ in COLUMNS])
        row = ws.max_row
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row, c)
            cell.fill = BANDS[band]
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        url_cell = ws.cell(row, 5)
        if url_cell.value:
            url_cell.hyperlink = url_cell.value
            url_cell.font = Font(color="0563C1", underline="single")
        for c in (7, 8):
            ws.cell(row, c).fill = YOUR_FILL
    dv = DataValidation(type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{ws.max_row}")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    help_ws = wb.create_sheet("Návod")
    for line in [
        "Ke každému řádku otevři odkaz ve sloupci E a porovnej hodnotu (D) se zdrojem. Citaci (F) najdeš přes Ctrl+F.",
        "Do žlutého sloupce G vyber verdikt:",
        "  ok = zdroj údaj potvrzuje",
        "  wrong = zdroj říká něco jiného (do H napiš co)",
        "  unverifiable = stránka nejde otevřít nebo tam údaj není",
        "  null_ok = údaj je prázdný a ani ty ho za 1–2 minuty nenajdeš",
        "  null_missed = údaj je prázdný, ale veřejně existuje (do H napiš kde)",
        "Částky: porovnej původní částku a měnu ve zdroji; přepočet do EUR dělá kód kurzem ECB.",
        "Investor se opakuje záměrně: každý řádek je jiný údaj (sloupec C).",
        "Sloupce I–L jsou jen pro kontext; verdikt AI (K) si nečti předem, ať tě neovlivní.",
        "Po vyplnění ulož a řekni Claudovi — verdikty načte: python src/manual_xlsx.py import",
    ]:
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 120
    wb.save(xlsx)
    print(f"uloženo: {xlsx} ({len(rows)} řádků)")


def import_(xlsx: Path) -> None:
    ws = load_workbook(xlsx, data_only=True)["Kontrola"]
    header = [c.value for c in ws[1]]
    col_of = {csv_col: header.index(h) for csv_col, h, _ in COLUMNS}
    filled = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        key = (row[col_of["id"]], row[col_of["field"]], row[col_of["group"]])
        filled[key] = (row[col_of["human_verdict"]] or "", row[col_of["human_comment"]] or "")
    with open(CSV_PATH, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    n = 0
    for r in rows:
        verdict, comment = filled.get((r["id"], r["field"], r["group"]), ("", ""))
        verdict = str(verdict).strip().lower()
        if verdict and verdict not in VERDICTS:
            raise ValueError(f"neplatný verdikt '{verdict}' u {r['id']}.{r['field']}")
        r["human_verdict"], r["human_comment"] = verdict, str(comment).strip()
        n += bool(verdict)
    with open(CSV_PATH, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    print(f"načteno {n}/{len(rows)} verdiktů do {CSV_PATH}")


if __name__ == "__main__":
    if len(sys.argv) < 2 or sys.argv[1] not in ("export", "import"):
        sys.exit(__doc__)
    path = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_XLSX
    export(path) if sys.argv[1] == "export" else import_(path)
